import docx
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_docx(output_path):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    COLOR_BLUE = RGBColor(37, 99, 235)       # #2563EB
    COLOR_DARK = RGBColor(15, 23, 42)        # #0F172A
    COLOR_MUTED = RGBColor(71, 85, 105)      # #475569

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # PAGE DE GARDE
    # -------------------------------------------------------------
    p_header1 = doc.add_paragraph()
    p_header1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_header1.add_run("UNIVERSITÉ HASSAN II DE CASABLANCA\n")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    
    r2 = p_header1.add_run("FACULTÉ DES SCIENCES AÏN CHOCK (FSAC)\n")
    r2.font.size = Pt(12)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_BLUE

    r3 = p_header1.add_run("Master Big Data & Cloud Computing (BD2C)\n")
    r3.font.size = Pt(11)
    r3.font.color.rgb = COLOR_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # Document Title Box Table
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = title_table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    cell.width = Inches(6.5)

    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(15)
    p_title.paragraph_format.space_after = Pt(15)

    rt1 = p_title.add_run("CAHIER DES CHARGES FONCTIONNEL ET TECHNIQUE\n\n")
    rt1.font.size = Pt(14)
    rt1.font.bold = True
    rt1.font.color.rgb = COLOR_BLUE

    rt2 = p_title.add_run("CONCEPTION ET DÉPLOIEMENT D’UNE PLATEFORME INTELLIGENTE DE DÉTECTION DE LA PNEUMONIE\n")
    rt2.font.size = Pt(15)
    rt2.font.bold = True
    rt2.font.color.rgb = COLOR_NAVY

    rt3 = p_title.add_run("\nÀ partir de radiographies thoraciques, basée sur le Deep Learning et une architecture MLOps")
    rt3.font.size = Pt(11)
    rt3.font.italic = True
    rt3.font.color.rgb = COLOR_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Metadata Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Nature du système :", "Prototype académique d'aide à la décision (Non certifié dispositif médical)"),
        ("Auteurs / Étudiants :", "OUAKIB Oussama & Mohammed Lemghari"),
        ("Encadrant Académique :", "M. Zouhair Chiba"),
        ("Modélisation Processus :", "Bizagi Modeler (Norme BPMN 2.0)"),
        ("Formation :", "Master Big Data & Cloud Computing (BD2C)"),
        ("Année Universitaire :", "2025 – 2026")
    ]
    for idx, (lbl, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.3)
        cell_val.width = Inches(4.2)
        
        p_lbl = cell_lbl.paragraphs[0]
        r_l = p_lbl.add_run(lbl)
        r_l.font.bold = True
        r_l.font.color.rgb = COLOR_NAVY
        
        p_val = cell_val.paragraphs[0]
        p_val.add_run(val)

    set_table_borders(meta_table, color="E2E8F0")

    doc.add_page_break()

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLUE
        return p

    def add_alert(title, text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "EFF6FF")
        cell.width = Inches(6.5)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        r_t = p.add_run(f"⚠️ {title}\n")
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_NAVY
        p.add_run(text)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # CONTENT
    add_h1("1. CONTEXTE ET PROBLÉMATIQUE")
    add_h2("1.1. Contexte Médical et Enjeu de Santé Publique")
    doc.add_paragraph(
        "La pneumonie est une infection respiratoire aiguë touchant le parenchyme pulmonaire. Selon l'OMS, elle constitue "
        "l'une des principales causes de mortalité infantile et gériatrique à l'échelle internationale."
    )
    doc.add_paragraph("La lecture des radiographies thoraciques comporte 3 contraintes cliniques majeures :")
    p_b1 = doc.add_paragraph(style='List Bullet')
    p_b1.add_run("Variabilité inter/intra-observateur : ").bold = True
    p_b1.add_run("Dépendance vis-à-vis de l'expérience du praticien et risque d'erreur en cas de fatigue.")

    p_b2 = doc.add_paragraph(style='List Bullet')
    p_b2.add_run("Complexité des lésions visuelles : ").bold = True
    p_b2.add_run("Infiltrats aux opacités subtiles et superpositions anatomiques (côtes, cœur).")

    p_b3 = doc.add_paragraph(style='List Bullet')
    p_b3.add_run("Engorgement des services : ").bold = True
    p_b3.add_run("Délais d'attente pouvant retarder la prise en charge thérapeutique.")

    add_h2("1.2. Problématique Technique et Scientifique")
    issues = [
        ("Fuite de données (Data Leakage) par récurrence patient :", " Isolation stricte des clichés d'un même patient au sein d'un seul sous-ensemble (Train/Val/Test)."),
        ("Opacité de l'IA (Boîte Noire) :", " Nécessité d'associer à chaque prédiction une carte d'explicabilité visuelle (Grad-CAM)."),
        ("Benchmarking d'architectures CNN :", " Protocole expérimental comparant Baseline CNN, ResNet50, DenseNet121 et EfficientNet-B0."),
        ("Calibration clinique du seuil :", " Priorité au Rappel (Recall >= 90%) pour minimiser les Faux Négatifs sanitaires."),
        ("Modélisation Bizagi BPMN & MLOps :", " Cartographie sous Bizagi Modeler et packaging microservices (API FastAPI, UI Streamlit, MLflow, MinIO, Docker, Kubernetes, Prometheus, Grafana).")
    ]
    for title, desc in issues:
        p_i = doc.add_paragraph(style='List Number')
        r_t = p_i.add_run(title)
        r_t.bold = True
        p_i.add_run(desc)

    add_alert("IMPORTANT : Statut Réglementaire", "Le présent système est un prototype académique d'aide à la décision. N'étant pas certifié comme dispositif médical, ses prédictions sont strictly destinées à fournir un second avis d'assistance au médecin praticien.")

    # -------------------------------------------------------------
    # BIZAGI MODELER / BPMN 2.0 SECTION
    # -------------------------------------------------------------
    add_h1("2. MODÉLISATION DES PROCESSUS MÉTIERS (BIZAGI MODELER / BPMN 2.0)")
    doc.add_paragraph(
        "Afin d'assurer une parfaite adéquation entre la conception technique et le workflow clinique, "
        "les processus de la plateforme ont été modélisés sous Bizagi Modeler conformément à la norme BPMN 2.0."
    )

    add_h2("2.1. Couloirs de Responsabilité (Swimlanes BPMN)")
    doc.add_paragraph("Le modèle définit 4 couloirs d'exécution principaux :")
    bpmn_lanes = [
        ("Médecin Radiologue / Praticien :", " Importe le cliché radiographique, sélectionne la sensibilité du seuil, examine la carte Grad-CAM et émet le diagnostic final."),
        ("Interface Applicative (Streamlit UI) :", " Assure la saisie des entrées, la communication REST et l'affichage dynamique des métriques visuelles."),
        ("Moteur d'Inférence & API (FastAPI) :", " Contrôle l'unicité SHA-256, exécute l'inférence CNN, génère la matrice Grad-CAM et transmet la réponse."),
        ("Infrastructure MLOps & Registre :", " Enregistre le modèle versionné dans MLflow/MinIO, persiste les historiques et exporte les métriques Prometheus.")
    ]
    for l_title, l_desc in bpmn_lanes:
        p_l = doc.add_paragraph(style='List Bullet')
        r_l = p_l.add_run(l_title)
        r_l.bold = True
        p_l.add_run(l_desc)

    add_h2("2.2. Matrice des Composants BPMN (Bizagi Modeler)")
    bpmn_table = doc.add_table(rows=1, cols=3)
    bpmn_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_hdr = bpmn_table.rows[0].cells
    b_titles = ["Élément BPMN", "Type / Symbole Bizagi", "Rôle dans le Système"]
    b_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]

    for i, t in enumerate(b_titles):
        b_hdr[i].text = t
        b_hdr[i].width = b_widths[i]
        set_cell_background(b_hdr[i], "1E3A8A")
        p = b_hdr[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    bpmn_data = [
        ("Événement de début", "Message Start Event", "Réception d'une radiographie sur l'interface clinique."),
        ("Tâche 1 (Automatisée)", "Service Task", "Contrôle de validité et calcul du hash SHA-256 anti-doublon."),
        ("Tâche 2 (Machine Learning)", "Script Task / ML Task", "Inférence CNN et calcul de la heatmap Grad-CAM."),
        ("Passerelle de décision", "Exclusive Gateway (XOR)", "Si Probabilité >= 67% -> Alerte Pneumonie ; Sinon -> Normal."),
        ("Tâche Utilisateur", "User Task", "Validation du compte rendu par le médecin radiologue."),
        ("Événement de fin", "End Event", "Archivage des métadonnées de prédiction et clôture du dossier.")
    ]

    for elem, typ, role in bpmn_data:
        row_cells = bpmn_table.add_row().cells
        row_cells[0].text = elem
        row_cells[1].text = typ
        row_cells[2].text = role
        row_cells[0].paragraphs[0].runs[0].font.bold = True

    set_table_borders(bpmn_table, color="CBD5E1")

    # -------------------------------------------------------------
    # SECTION 3: BESOINS
    # -------------------------------------------------------------
    add_h1("3. SPÉCIFICATION DES BESOINS")
    add_h2("3.1. Besoins Fonctionnels (RF)")

    col_widths = [Inches(1.0), Inches(2.2), Inches(3.3)]
    rf_table = doc.add_table(rows=1, cols=3)
    rf_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = rf_table.rows[0].cells
    hdr_titles = ["Code", "Exigence Fonctionnelle", "Description Détaillée"]
    
    for i, t in enumerate(hdr_titles):
        hdr_cells[i].text = t
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A")
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    rf_data = [
        ("RF-01", "Ingestion & Dédoublonnage", "Validation du format des images, élimination des doublons exacts par hashage SHA-256 et enregistrement du catalogue."),
        ("RF-02", "Partitionnement Patient", "Split stratifié garantissant l'étanchéité stricte des patients entre les ensembles Train, Validation et Test."),
        ("RF-03", "Gestion MLOps", "Traçabilité et versioning automatique des hyperparamètres, métriques et modèles dans MLflow, MinIO et PostgreSQL."),
        ("RF-04", "Calibration du Seuil", "Calcul automatisé du seuil optimal de décision sur le jeu de validation sous contrainte de rappel minimum (>= 90%)."),
        ("RF-05", "Explicabilité Visuelle", "Génération de cartes d'attention Grad-CAM superposées à la radiographie pour surligner les foyers de condensation."),
        ("RF-06", "Service API REST", "API FastAPI exposant les endpoints /predict (inférence + Grad-CAM), /health et /metrics."),
        ("RF-07", "Interface UI Clinique", "Dashboard Streamlit interactif permettant l'import de clichés, la simulation de seuils et l'affichage des résultats."),
        ("RF-08", "Supervision & Alerting", "Exportation des métriques vers Prometheus et visualisation sur tableaux de bord Grafana avec alertes automatiques.")
    ]

    for code, title, desc in rf_data:
        row_cells = rf_table.add_row().cells
        row_cells[0].text = code
        row_cells[1].text = title
        row_cells[2].text = desc
        row_cells[0].paragraphs[0].runs[0].font.bold = True

    set_table_borders(rf_table, color="CBD5E1")

    # -------------------------------------------------------------
    # SECTION 4: PLANNING
    # -------------------------------------------------------------
    add_h1("4. PLANIFICATION DU PROJET (104 HEURES)")
    plan_table = doc.add_table(rows=1, cols=4)
    plan_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_hdr = plan_table.rows[0].cells
    p_titles = ["Phase", "Intitulé de la Phase", "Durée", "Livrables Attendus"]
    p_widths = [Inches(0.9), Inches(2.2), Inches(0.8), Inches(2.6)]

    for i, t in enumerate(p_titles):
        p_hdr[i].text = t
        p_hdr[i].width = p_widths[i]
        set_cell_background(p_hdr[i], "1E3A8A")
        p = p_hdr[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    plan_data = [
        ("Phase 1", "Cadrage & Modélisation Bizagi BPMN", "10h", "Cahier des charges validé, diagrammes BPMN 2.0."),
        ("Phase 2", "Infrastructure MLOps Locale", "12h", "Services PostgreSQL, MinIO, MLflow conteneurisés."),
        ("Phase 3", "Data Pipeline & Split Patient", "14h", "Catalogue d'images, dédoublonnage SHA-256, splits étanches."),
        ("Phase 4", "Baseline CNN PyTorch", "12h", "Modèle baseline entraîné et suivi dans MLflow."),
        ("Phase 5", "Benchmark Architectures Avancées", "16h", "Modèles ResNet50, DenseNet121, EfficientNet-B0 comparés."),
        ("Phase 6", "Calibration & Explicabilité Grad-CAM", "12h", "Seuil calibré à 0.67 (Rappel >= 90%), module Grad-CAM."),
        ("Phase 7", "API FastAPI & Interface Streamlit", "12h", "Endpoints API opérationnels, dashboard UI interactif."),
        ("Phase 8", "CI/CD GitHub Actions & Registre GHCR", "8h", "Workflows CI/CD, images Docker optimisées non-root."),
        ("Phase 9", "Déploiement Kubernetes & Grafana", "8h", "Manifestes K8s, tableaux Grafana, rapport final PFE.")
    ]

    for ph, title, dur, liv in plan_data:
        row_cells = plan_table.add_row().cells
        row_cells[0].text = ph
        row_cells[1].text = title
        row_cells[2].text = dur
        row_cells[3].text = liv
        row_cells[0].paragraphs[0].runs[0].font.bold = True

    set_table_borders(plan_table, color="CBD5E1")

    # -------------------------------------------------------------
    # SECTION 5: RECETTE & RESULTATS
    # -------------------------------------------------------------
    add_h1("5. MATRICE DE RECETTE ET RÉSULTATS EXPÉRIMENTAUX")
    rec_table = doc.add_table(rows=1, cols=4)
    rec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_hdr = rec_table.rows[0].cells
    r_titles = ["Métrique Académique", "Valeur Obtenue (Test)", "Objectif du CDC", "Statut"]
    r_widths = [Inches(2.5), Inches(1.3), Inches(1.3), Inches(1.4)]

    for i, t in enumerate(r_titles):
        r_hdr[i].text = t
        r_hdr[i].width = r_widths[i]
        set_cell_background(r_hdr[i], "1E3A8A")
        p = r_hdr[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    rec_data = [
        ("Accuracy Global", "86,72 %", ">= 85,00 %", "CONFORME"),
        ("Précision (Classe Pneumonie)", "91,00 %", ">= 88,00 %", "CONFORME"),
        ("Rappel (Recall Classe Pneumonie)", "90,55 %", ">= 90,00 %", "CONFORME"),
        ("Spécificité (Classe Normal)", "76,82 %", ">= 70,00 %", "CONFORME"),
        ("F1-Score (Classe Pneumonie)", "90,77 %", ">= 88,00 %", "CONFORME"),
        ("ROC-AUC", "94,05 %", ">= 92,00 %", "CONFORME"),
        ("PR-AUC", "97,71 %", ">= 95,00 %", "CONFORME"),
        ("Taux de Faux Négatifs", "9,45 %", "<= 10,00 %", "CONFORME")
    ]

    for m, val, obj, st in rec_data:
        row_cells = rec_table.add_row().cells
        row_cells[0].text = m
        row_cells[1].text = val
        row_cells[2].text = obj
        row_cells[3].text = st
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[3].paragraphs[0].runs[0].font.bold = True
        row_cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 101, 52)

    set_table_borders(rec_table, color="CBD5E1")

    # -------------------------------------------------------------
    # SECTION 6: SIGNATURES
    # -------------------------------------------------------------
    add_h1("6. SIGNATURES ET VALIDATION")
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_cells = sig_table.rows[0].cells
    sig_names = [
        ("Étudiant 1 :", "OUAKIB Oussama"),
        ("Étudiant 2 :", "Mohammed Lemghari"),
        ("Encadrant Académique :", "M. Zouhair Chiba")
    ]
    for idx, (role, name) in enumerate(sig_names):
        cell = sig_cells[idx]
        cell.width = Inches(2.1)
        set_cell_background(cell, "F8FAFC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        r_r = p.add_run(f"{role}\n")
        r_r.font.bold = True
        r_r.font.color.rgb = COLOR_NAVY
        p.add_run(f"{name}\n\nDate : ____/____/2026\n\nSignature :")

    set_table_borders(sig_table, color="CBD5E1")

    # Save Document safely
    try:
        doc.save(output_path)
        print(f"Successfully generated docx: {output_path}")
    except PermissionError:
        alt_path = output_path.replace(".docx", "_v2.docx")
        doc.save(alt_path)
        print(f"Original locked. Successfully generated docx to alternative path: {alt_path}")

if __name__ == "__main__":
    build_docx("c:\\Users\\GALTOUT\\Desktop\\PFE2-main\\PFE2-main\\reports\\Cahier_des_Charges_PFE.docx")
